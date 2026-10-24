"""Authenticated XLSX round trips against synthetic SQLite, never production."""
import asyncio
from dataclasses import replace
from datetime import datetime, timedelta
import hashlib
import logging
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, inspect
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.admin import COOKIE, credential_key, now
from app.core.admin_security import admin_security
from app.core.admin_logging import AdminAccessFilter
from app.core.config import Settings, settings
from app.core.database import Base, get_db
from app.main import app
from app.models.db_models import ContactSubmission
from app.services import export_service as exports

URL = '/api/admin/submissions/export'


class ExportTests(unittest.TestCase):
    def setUp(self):
        admin_security.clear()
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        def test_db():
            with self.sessions() as db:
                yield db
        app.dependency_overrides[get_db] = test_db
        self.credentials = patch.object(Settings, 'admin_credentials', return_value=('test-admin', 'test-password'))
        self.credentials.start()
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name) / 'exports'
        self.root_patch = patch.object(exports, 'EXPORT_ROOT', self.root)
        self.root_patch.start()
        self.client = TestClient(app, base_url='https://testserver', headers={'X-Admin-Request': '1'})
        with self.sessions() as db:
            for index in range(25):
                db.add(ContactSubmission(name=f'Person {index}', email=f'person{index}@example.org',
                                         phone_number='001234567890', message='First line\nSecond line',
                                         created_at=datetime(2026, 10, 20) + timedelta(days=index // 10)))
            db.commit()

    def tearDown(self):
        self.client.close()
        self.credentials.stop()
        self.root_patch.stop()
        self.temp.cleanup()
        credential_key.cache_clear()
        app.dependency_overrides.clear()
        self.engine.dispose()
        admin_security.clear()

    def login(self):
        self.assertEqual(self.client.post('/api/admin/login', json={
            'username': 'test-admin', 'password': 'test-password'}).status_code, 200)

    def workbook(self, **params):
        response = self.client.get(URL, params=params)
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
        self.assertEqual(response.headers['content-type'], exports.MEDIA_TYPE)
        self.assertRegex(response.headers['content-disposition'], r'attachment; filename="neasw-submissions-\d{4}-\d\d-\d\d-\d{6}-UTC.xlsx"')
        for name, value in (('cache-control', 'no-store'), ('pragma', 'no-cache'), ('x-content-type-options', 'nosniff')):
            self.assertEqual(response.headers[name], value)
        self.assertEqual(list(self.root.iterdir()), [])
        return load_workbook(BytesIO(response.content))

    def test_auth_guard_expiry_and_logout(self):
        self.assertEqual(self.client.get(URL).status_code, 401)
        self.client.cookies.set(COOKIE, 'forged')
        self.assertEqual(self.client.get(URL).status_code, 401)
        self.client.cookies.clear()
        self.login()
        self.assertEqual(self.client.get(URL, headers={'X-Admin-Request': ''}).status_code, 403)
        self.assertEqual(self.client.get(URL, headers={'Origin': 'https://evil.example'}).status_code, 403)
        token_hash = hashlib.sha256(self.client.cookies.get(COOKIE).encode()).hexdigest()
        session = admin_security.get_session(token_hash, now())
        admin_security.create_session(replace(session, expires_at=now() - timedelta(seconds=1)), now())
        self.assertEqual(self.client.get(URL).status_code, 401)
        self.login()
        token = self.client.cookies.get(COOKIE)
        self.client.post('/api/admin/logout')
        self.client.cookies.set(COOKIE, token)
        self.assertEqual(self.client.get(URL).status_code, 401)

    def test_all_pages_stable_order_and_read_only(self):
        self.login()
        statements = []
        def record(connection, cursor, statement, parameters, context, many):
            statements.append(statement)
        event.listen(self.engine, 'before_cursor_execute', record)
        try:
            workbook = self.workbook()
        finally:
            event.remove(self.engine, 'before_cursor_execute', record)
        sheet = workbook['Submissions']
        self.assertEqual(sheet.max_row, 26)
        self.assertEqual([row[0].value for row in list(sheet.rows)[1:]], list(range(25, 0, -1)))
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:H26')
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet['D2'].value, '001234567890')
        self.assertEqual(sheet['D2'].data_type, 's')
        self.assertIsNone(sheet['E2'].value)
        self.assertIsNone(sheet['F2'].value)
        self.assertEqual(sheet['G2'].value, 'First line\nSecond line')
        self.assertEqual(sheet['H2'].value, datetime(2026, 10, 22))
        self.assertIn('UTC', sheet['H2'].number_format)
        self.assertEqual(dict(workbook['Export info'].values)['Row count'], 25)
        self.assertTrue(all(sql.lstrip().upper().startswith(('BEGIN', 'SELECT')) for sql in statements), statements)
        self.assertEqual(inspect(self.engine).get_table_names(), ['contact_submissions'])
        with self.sessions() as db:
            self.assertEqual(db.query(ContactSubmission).count(), 25)
        workbook.close()

    def test_matching_search_dates_and_empty_results(self):
        self.login()
        for params in ({'date_from': '2026-10-21', 'date_to': '2026-10-21'},
                       {'search': 'person2', 'date_to': '2026-10-20'},
                       {'search': 'Second line', 'date_from': '2026-10-21'}, {'search': 'no matching record'}):
            with self.subTest(params=params):
                listing = self.client.get('/api/admin/submissions', params={**params, 'page_size': 100}).json()
                workbook = self.workbook(**params)
                rows = list(workbook['Submissions'].values)[1:]
                self.assertEqual([row[0] for row in rows], [row['id'] for row in listing['items']])
                self.assertEqual(len(rows), listing['total'])
                self.assertEqual(dict(workbook['Export info'].values)['Row count'], listing['total'])
                workbook.close()
        self.assertEqual(self.client.get('/api/admin/submissions', params={'date_from': '2026-10-21', 'date_to': '2026-10-21'}).json()['total'], 10)
        for params in ({'date_from': 'bad'}, {'date_from': '2026-10-22', 'date_to': '2026-10-20'},
                       {'date_to': '9999-12-31'}, {'search': 'x' * 201}):
            for url in (URL, '/api/admin/submissions'):
                self.assertEqual(self.client.get(url, params=params).status_code, 422)

    def test_literal_strings_international_text_and_long_fields(self):
        self.login()
        long_value = '=untrusted ' + ('अ😀\n' * 15000)
        with self.sessions() as db:
            record = db.get(ContactSubmission, 25)
            record.name = '=HYPERLINK("https://evil.example")'
            record.email = '@unsafe@example.org'
            record.phone_number = '+001234567'
            record.organization_name = '-1+2'
            record.preferred_date = 'অসমীয়া 日本語'
            record.message = long_value
            db.commit()
        workbook = self.workbook(search='unsafe')
        sheet = workbook['Submissions']
        for address, expected in (('B2', '=HYPERLINK("https://evil.example")'), ('C2', '@unsafe@example.org'),
                                  ('D2', '+001234567'), ('E2', '-1+2'), ('F2', 'অসমীয়া 日本語')):
            self.assertEqual(sheet[address].value, expected)
            self.assertEqual(sheet[address].data_type, 's')
            self.assertIsNone(sheet[address].hyperlink)
        self.assertIn('See Long text: submission 25, field Message', sheet['G2'].value)
        chunks = list(workbook['Long text'].values)[1:]
        self.assertEqual(''.join(row[3] for row in chunks), long_value)
        self.assertEqual([row[2] for row in chunks], list(range(1, len(chunks) + 1)))
        self.assertTrue(all(len(row[3].encode('utf-16-le')) // 2 <= 32767 for row in chunks))
        self.assertTrue(all(cell.data_type != 'f' and not cell.hyperlink for sheet in workbook for row in sheet for cell in row))
        workbook.close()

    def test_xml_incompatible_characters_are_explicit(self):
        self.login()
        with self.sessions() as db:
            db.get(ContactSubmission, 25).message = 'control\x00text\x0b'
            db.commit()
        workbook = self.workbook()
        self.assertEqual(workbook['Submissions']['G2'].value, 'control\\u0000text\\u000B')
        workbook.close()

    def test_limits_concurrency_and_failures_cleanup(self):
        self.login()
        with patch.object(settings, 'EXPORT_MAX_ROWS', 2):
            response = self.client.get(URL)
            self.assertEqual(response.status_code, 422)
            self.assertIn('Narrow', response.json()['detail'])
        with patch.object(settings, 'EXPORT_MAX_FILE_MB', 0):
            self.assertEqual(self.client.get(URL).status_code, 422)
        self.assertTrue(exports.EXPORT_SLOTS.acquire(blocking=False))
        try:
            response = self.client.get(URL)
            self.assertEqual(response.status_code, 429)
            self.assertEqual(response.headers['retry-after'], '30')
        finally:
            exports.EXPORT_SLOTS.release()
        with patch.object(exports.Workbook, 'save', side_effect=OSError('synthetic disk failure')):
            self.assertEqual(self.client.get(URL).status_code, 500)
        self.assertEqual(list(self.root.iterdir()), [])
        self.assertFalse(any(str(self.root) in path for path in exports.ALL_TEMP_FILES))
        self.workbook().close()  # failed attempts released the slot
        # Exercise the completed ZIP cap independently of the source text cap.
        with patch.object(settings, 'EXPORT_MAX_FILE_MB', 0.001):
            response = self.client.get(URL, params={'search': 'no matches'})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(list(self.root.iterdir()), [])
        self.workbook().close()

    def test_source_budget_validation_and_private_access_logs(self):
        self.login()
        with patch.object(settings, 'EXPORT_MAX_FILE_MB', 1):
            with self.sessions() as db:
                db.get(ContactSubmission, 25).message = 'x' * (1024 * 1024 + 1)
                db.commit()
            response = self.client.get(URL)
            self.assertEqual(response.status_code, 422)
            self.assertIn('too much text', response.json()['detail'])
        for values in ({'EXPORT_MAX_ROWS': 0}, {'EXPORT_MAX_FILE_MB': 101}, {'EXPORT_MAX_CONCURRENT': 3}):
            with self.assertRaises(ValueError):
                Settings(_env_file=None, **values)
        record = logging.LogRecord('uvicorn.access', logging.INFO, '', 0, '%s %s %s %s %s',
                                   ('127.0.0.1', 'GET', '/api/admin/submissions/export?search=private', '1.1', 200), None)
        AdminAccessFilter().filter(record)
        self.assertNotIn('private', record.getMessage())
        self.assertIn('/api/admin/submissions/export', record.getMessage())
        self.assertEqual(list(self.root.iterdir()), [])

    def test_disconnect_and_stale_directory_cleanup(self):
        with self.sessions() as db:
            artifact = exports.generate_export(db)
        response = exports.ExportResponse(artifact)
        async def send(message):
            raise OSError('synthetic disconnected client')
        async def receive():
            return {'type': 'http.disconnect'}
        with self.assertRaises(OSError):
            asyncio.run(response({'type': 'http', 'method': 'GET', 'headers': [],
                                  'asgi': {'spec_version': '2.4'}}, receive, send))
        self.assertFalse(artifact.directory.exists())
        self.assertTrue(exports.EXPORT_SLOTS.acquire(blocking=False))
        exports.EXPORT_SLOTS.release()
        stale = self.root / 'export-stale'
        stale.mkdir()
        (stale / 'leftover.xml').write_text('private synthetic data')
        os.utime(stale, (0, 0))
        other = self.root / 'unrelated'
        other.mkdir()
        exports.clean_stale_exports()
        self.assertFalse(stale.exists())
        self.assertTrue(other.exists())
        with self.sessions() as db:
            active = exports.generate_export(db)
        try:
            os.utime(active.directory, (0, 0))
            exports.clean_stale_exports()
            self.assertTrue(active.path.exists())
        finally:
            active.cleanup()


if __name__ == '__main__':
    unittest.main()
