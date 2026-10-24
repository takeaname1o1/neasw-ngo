"""Bounded, read-only XLSX exports; no submission content is logged."""
from dataclasses import dataclass
from datetime import date, datetime, timezone
import fcntl
from itertools import chain
import logging
import os
from pathlib import Path
import re
import shutil
from tempfile import gettempdir, mkdtemp
from threading import BoundedSemaphore
from time import monotonic, time

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import ALL_TEMP_FILES, WorksheetWriter
from sqlalchemy import func, text
from sqlalchemy.orm import Session
from starlette.responses import FileResponse

from app.core.config import settings
from app.models.db_models import ContactSubmission
from app.services.submission_queries import submission_query

logger = logging.getLogger(__name__)
EXPORT_ROOT = Path(gettempdir()) / 'neasw-private-exports'
EXPORT_SLOTS = BoundedSemaphore(settings.EXPORT_MAX_CONCURRENT)
MEDIA_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HEADERS = ('Submission ID', 'Name', 'Email', 'Phone number', 'Organisation / plan',
           'Preferred date', 'Message', 'Submitted at (UTC)')
COLUMNS = (ContactSubmission.id, ContactSubmission.name, ContactSubmission.email,
           ContactSubmission.phone_number, ContactSubmission.organization_name,
           ContactSubmission.preferred_date, ContactSubmission.message, ContactSubmission.created_at)
# XML 1.0 cannot represent these characters. Display explicit escapes rather than lose them.
INVALID_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def clean_stale_exports():
    """Reap crashed exports after one hour; active exports in this process are excluded."""
    if not EXPORT_ROOT.exists():
        return
    for directory in EXPORT_ROOT.glob('export-*'):
        try:
            if (directory.is_symlink() or not directory.is_dir()
                    or time() - directory.stat().st_mtime <= 3600):
                continue
            # An advisory flock also protects an unusually long-running active response.
            with (directory / '.lock').open('a') as lock:
                fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
                shutil.rmtree(directory)
        except (BlockingIOError, FileNotFoundError):
            continue
        except OSError as error:
            logger.warning('Export stale cleanup error=%s', type(error).__name__)


@dataclass
class ExportArtifact:
    directory: Path
    path: Path
    filename: str
    row_count: int
    lock: object
    closed: bool = False

    def cleanup(self):
        if not self.closed:
            self.closed = True
            try:
                shutil.rmtree(self.directory, ignore_errors=True)
            finally:
                self.lock.close()
                EXPORT_SLOTS.release()


class ExportResponse(FileResponse):
    def __init__(self, artifact: ExportArtifact):
        super().__init__(artifact.path, media_type=MEDIA_TYPE, filename=artifact.filename)
        self.artifact = artifact

    async def __call__(self, scope, receive, send):
        try:
            await super().__call__(scope, receive, send)
        finally:
            # BackgroundTask alone does not cover a failed/disconnected send.
            self.artifact.cleanup()


def _sheet(workbook, directory, title, headers, widths):
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = 'A2'
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    # openpyxl 3.1's write-only writer otherwise uses global /tmp XML files.
    # Keep its intermediate XML in our private, crash-reapable directory too.
    output = str(directory / f'{len(workbook.worksheets)}.xml')
    sheet._writer = WorksheetWriter(sheet, out=output)
    ALL_TEMP_FILES.append(output)
    sheet._writer.write_top()
    cells = []
    for heading in headers:
        cell = _literal(sheet, heading)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0F766E')
        cells.append(cell)
    sheet.append(cells)
    return sheet


def _literal(sheet, value):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = 's'  # override openpyxl's automatic '=' formula detection
        cell.number_format = '@'
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def _chunks(value):
    """Excel counts UTF-16 units, including two units for supplementary characters."""
    start = units = 0
    for index, char in enumerate(value):
        size = 2 if ord(char) > 0xFFFF else 1
        if units + size > 32767:
            yield value[start:index]
            start, units = index, 0
        units += size
    yield value[start:]


def _close_workbook(workbook):
    for sheet in workbook.worksheets:
        try:
            if not sheet.closed:
                sheet.close()
        except Exception as error:
            logger.warning('Export writer cleanup error=%s', type(error).__name__)
        finally:
            writer = getattr(sheet, '_writer', None)
            if writer and writer.out in ALL_TEMP_FILES:
                ALL_TEMP_FILES.remove(writer.out)
    workbook.close()


def generate_export(db: Session, search: str = '', date_from: date | None = None,
                    date_to: date | None = None) -> ExportArtifact:
    if not EXPORT_SLOTS.acquire(blocking=False):
        raise HTTPException(429, 'Another export is in progress. Please try again shortly',
                            headers={'Retry-After': '30'})
    started = monotonic()
    artifact = None
    directory = lock = None
    workbook = None
    rows = None
    count = 0
    try:
        EXPORT_ROOT.mkdir(mode=0o700, parents=True, exist_ok=True)
        os.chmod(EXPORT_ROOT, 0o700)
        clean_stale_exports()
        directory = Path(mkdtemp(prefix='export-', dir=EXPORT_ROOT))
        lock = (directory / '.lock').open('a')
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        generated = datetime.now(timezone.utc)
        artifact = ExportArtifact(directory, directory / 'submissions.xlsx',
                                  generated.strftime('neasw-submissions-%Y-%m-%d-%H%M%S-UTC.xlsx'), 0, lock)
        # A single snapshot covers counts, text-size checks and selected rows.
        if db.bind.dialect.name == 'postgresql':
            db.connection(execution_options={'isolation_level': 'REPEATABLE READ'})
            db.execute(text('SET TRANSACTION READ ONLY'))
        elif db.bind.dialect.name == 'sqlite':
            db.connection().exec_driver_sql('BEGIN')
        query = submission_query(db, search, date_from, date_to)
        count = query.order_by(None).count()
        if count > settings.EXPORT_MAX_ROWS:
            raise HTTPException(422, f'Export exceeds {settings.EXPORT_MAX_ROWS:,} rows. Narrow the search or UTC date range')
        # Bound source text as well as compressed output, before fetching large fields.
        size = query.order_by(None).with_entities(func.sum(sum(
            func.coalesce(func.length(column), 0) for column in COLUMNS[1:7]))).scalar() or 0
        limit = settings.EXPORT_MAX_FILE_MB * 1024 * 1024
        if size > limit:
            raise HTTPException(422, 'Export contains too much text. Narrow the search or UTC date range')
        workbook = Workbook(write_only=True)
        main = _sheet(workbook, directory, 'Submissions', HEADERS, (16, 26, 34, 22, 32, 22, 72, 25))
        main.auto_filter.ref = f'A1:H{count + 1}'
        continuation = None
        rows = db.execute(query.with_entities(*COLUMNS).statement, execution_options={'yield_per': 500})
        for row_number, row in enumerate(rows, 2):
            cells = []
            for index, value in enumerate(row):
                if isinstance(value, str):
                    value = INVALID_XML.sub(lambda match: f'\\u{ord(match[0]):04X}', value)
                    chunks = iter(_chunks(value))
                    first = next(chunks)
                    second = next(chunks, None)
                    if second is not None:
                        if continuation is None:
                            continuation = _sheet(workbook, directory, 'Long text',
                                                  ('Submission ID', 'Field', 'Chunk', 'Text'), (16, 26, 12, 100))
                        for chunk_number, chunk in enumerate(chain((first, second), chunks), 1):
                            continuation.append([_literal(continuation, item) for item in
                                                 (row.id, HEADERS[index], chunk_number, chunk)])
                        value = f'See Long text: submission {row.id}, field {HEADERS[index]}'
                if isinstance(value, datetime):
                    value = value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
                cell = _literal(main, value)
                if index == 7 and value is not None:
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss "UTC"'
                cells.append(cell)
            main.row_dimensions[row_number].height = 48
            main.append(cells)
        info = _sheet(workbook, directory, 'Export info', ('Setting', 'Value'), (30, 100))
        for key, value in (
            ('Generated at (UTC)', generated.replace(tzinfo=None)), ('Row count', count),
            ('Search', search.strip() or '(all)'), ('From date (UTC, inclusive)', str(date_from) if date_from else '(any)'),
            ('To date (UTC, inclusive)', str(date_to) if date_to else '(any)'), ('Timezone', 'UTC'),
            ('Scope', 'All matching submissions across every dashboard page'),
            ('Long text', 'Values above 32,767 UTF-16 units are split into ordered chunks on Long text'),
            ('XML characters', 'XML-incompatible control characters are displayed as literal \\uNNNN escapes'),
        ):
            info.append([_literal(info, key), _literal(info, value)])
        workbook.save(artifact.path)
        os.chmod(artifact.path, 0o600)
        if artifact.path.stat().st_size > limit:
            raise HTTPException(422, 'Excel file exceeds the export size limit. Narrow the search or UTC date range')
        artifact.row_count = count
        logger.info('Export outcome=ready rows=%d duration=%.2fs', count, monotonic() - started)
        return artifact
    except Exception as error:
        logger.warning('Export outcome=failed rows=%d duration=%.2fs error=%s',
                       count, monotonic() - started, type(error).__name__)
        if workbook:
            _close_workbook(workbook)
            workbook = None
        if artifact:
            artifact.cleanup()
        else:
            if directory:
                shutil.rmtree(directory, ignore_errors=True)
            if lock:
                lock.close()
            EXPORT_SLOTS.release()
        if isinstance(error, HTTPException):
            raise
        raise HTTPException(500, 'Unable to prepare Excel export. Please try again') from error
    finally:
        try:
            if rows is not None:
                rows.close()
            if workbook:
                _close_workbook(workbook)
        finally:
            db.rollback()
